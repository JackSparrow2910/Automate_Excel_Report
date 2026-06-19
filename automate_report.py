from pivot_table import create_pivot_table
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font

import os
import sys

#Preparation to convert to exe and find the path of the exe file
application_path = os.path.dirname(sys.executable)

#Create the pivot_table from module pivot_table.py
month=input("Input month: ")
path=""
path=create_pivot_table(month,application_path)

if path=="":
    sys.exit()

#Configuration
workbook = load_workbook(path)
sheet = workbook['Sheet1']

#Active table
min_column = workbook.active.min_column
max_column = workbook.active.max_column
min_row = workbook.active.min_row
max_row = workbook.active.max_row

#Add barchart
barchart=BarChart()

data = Reference(sheet,min_col=min_column+1,min_row=min_row,max_col=max_column,max_row=max_row)
categories = Reference(sheet,min_col=min_column,min_row=min_row+1,max_col=min_column,max_row=max_row)

barchart.add_data(data,titles_from_data=True)
barchart.set_categories(categories)

barchart.x_axis.delete=False

sheet.add_chart(barchart, "B12")

barchart.title="City by Gender"
barchart.style=4

#Add headings of sheet
sheet["A1"]="Accommodation Report"
sheet["A1"].font = Font(bold=True,size=20)
sheet["A2"]=month
sheet["A2"].font = Font(bold=True,size=10)

#Evaluate sum of male and female citizens
sheet[f"{get_column_letter(min_column)}{max_row+1}"]="Sum"
for i in range(min_column+1,max_column+1):
    letter = get_column_letter(i)
    sheet[f"{letter}{max_row+1}"]=f"=SUM({letter}{min_row+1}:{letter}{max_row})"


#Save
workbook.save(path)
workbook.close()
